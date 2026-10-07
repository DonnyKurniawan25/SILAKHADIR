"""Strict, bounded XLSX attendance interchange. No certificate side effects."""
from collections import defaultdict
from io import BytesIO
import re
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.datavalidation import DataValidation

from apps.participants.models import Participant
from .models import Attendance

HEADERS = ('NIK', 'NIP', 'Nama Lengkap', 'Instansi', 'Jabatan', 'No HP', 'Email', 'Status')
FIELDS = ('nik', 'nip', 'full_name', 'institution', 'position', 'phone', 'email', 'status')
MAX_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_ROWS = 5000
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class XlsxError(ValueError):
    pass


def _text_cell(cell, value):
    # Explicit string type (not apostrophe escaping): even =,+,-,@ remain inert
    # text in Excel, without changing the exported user's actual value.
    cell.value = str(value or '')
    cell.data_type = 's'
    cell.number_format = '@'
    cell.font = Font(name='Arial', size=11)


def _sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='245A81')
    for values in rows:
        index = ws.max_row + 1
        for col, value in enumerate(values, 1):
            _text_cell(ws.cell(index, col), value)
    for col, width in zip('ABCDEFGH', (22, 24, 32, 32, 26, 22, 32, 20)):
        ws.column_dimensions[col].width = width
        ws.column_dimensions[col].number_format = '@'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def workbook_bytes(event, template=False):
    wb = Workbook()
    wb.remove(wb.active)
    rows = [] if template else (
        [*(getattr(a.participant, f) for f in FIELDS[:-1]), a.status]
        for a in Attendance.objects.filter(event=event).select_related('participant').order_by('participant__full_name', 'participant__nik')
    )
    ws = _sheet(wb, 'Absensi', rows)
    if template:
        status = DataValidation(type='list', formula1='"hadir,tidak_hadir"')
        status.errorTitle = 'Status tidak valid'
        status.error = 'Pilih hadir atau tidak_hadir.'
        status.showErrorMessage = True
        ws.add_data_validation(status)
        status.add(f'H2:H{MAX_ROWS + 1}')
        instructions = wb.create_sheet('Petunjuk')
        for text in (
            'Isi hanya sheet Absensi; sheet Contoh tidak diimpor.',
            'Jangan ubah nama/urutan delapan kolom pada baris pertama.',
            'NIK wajib 16 digit; NIP opsional 18 digit. Simpan NIK, NIP, No HP sebagai TEXT, bukan angka.',
            'Nama Lengkap wajib. Status wajib hadir atau tidak_hadir.',
            'Batas 5000 baris data dan ukuran file 5 MB. Rumus/formula tidak diperbolehkan.',
            'NIK/nama duplikat atau identitas ambigu ditolak. Kolom opsional kosong tidak menghapus data lama.',
            'Preview tidak menyimpan data. Konfirmasikan dry_run=false untuk menyimpan seluruh baris valid secara atomik.',
        ):
            _text_cell(instructions.cell(instructions.max_row + (1 if instructions['A1'].value else 0), 1), text)
        instructions.column_dimensions['A'].width = 120
        _sheet(wb, 'Contoh', [['0123456789012345', '012345678901234567', 'Nama Contoh', 'Diskominfo', 'Staf', '08123456789', 'contoh@example.com', 'hadir']])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _check_archive(content):
    """Bound resources before openpyxl parses shared strings/styles/dimensions."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 100 or len({m.filename for m in members}) != len(members):
                raise XlsxError('Arsip XLSX memiliki terlalu banyak/duplikat komponen.')
            if sum(m.file_size for m in members) > MAX_EXPANDED_BYTES:
                raise XlsxError('Ukuran hasil ekstraksi XLSX melebihi 25 MB.')
            for member in members:
                if member.flag_bits & 1 or member.filename.startswith('/') or '..' in member.filename.split('/'):
                    raise XlsxError('Komponen arsip XLSX tidak aman.')
                if 'vbaProject' in member.filename or member.filename.startswith('xl/externalLinks/'):
                    raise XlsxError('Macro dan tautan workbook eksternal tidak diperbolehkan.')
                if member.file_size > 1024 * 1024 and member.file_size / max(member.compress_size, 1) > 200:
                    raise XlsxError('Rasio kompresi XLSX tidak aman.')
                if not member.filename.endswith('.xml'):
                    continue
                xml = archive.read(member)
                markup = xml.replace(b'\x00', b'').upper()  # also UTF-16/32 XML
                if b'<!DOCTYPE' in markup or b'<!ENTITY' in markup:
                    raise XlsxError('Deklarasi DTD/entity tidak diperbolehkan.')
                # Scan every worksheet, including ignored example/instruction sheets.
                if not member.filename.startswith('xl/worksheets/'):
                    continue
                for _, element in ElementTree.iterparse(BytesIO(xml), events=('end',)):
                    tag = element.tag.rsplit('}', 1)[-1]
                    if tag == 'f':
                        raise XlsxError('Formula tidak diperbolehkan dalam workbook.')
                    if tag == 'row' and (int(element.get('r', '0')) > MAX_ROWS + 1):
                        raise XlsxError('Maksimal 5000 baris data per sheet.')
                    if tag == 'c':
                        ref = re.fullmatch(r'([A-Z]+)([0-9]+)', element.get('r', ''))
                        if not ref or column_index_from_string(ref[1]) > 64 or int(ref[2]) > MAX_ROWS + 1:
                            raise XlsxError('Rentang sel XLSX terlalu besar/tidak valid.')
                    element.clear()
    except XlsxError:
        raise
    except (BadZipFile, ElementTree.ParseError, KeyError, RuntimeError, OverflowError, ValueError, NotImplementedError) as exc:
        raise XlsxError('File bukan workbook XLSX yang valid.') from exc


def read_rows(upload):
    if not upload or not upload.name.lower().endswith('.xlsx'):
        raise XlsxError('Unggah file .xlsx pada field file.')
    if upload.size > MAX_BYTES:
        raise XlsxError('Ukuran file maksimal 5 MB.')
    content = upload.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise XlsxError('Ukuran file maksimal 5 MB.')
    _check_archive(content)
    wb = None
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(wb.worksheets) > 10:
            raise XlsxError('Maksimal 10 sheet dalam workbook.')
        # Excel's first sheet is accepted for third-party workbooks, except the
        # known help/example tabs which must never become import data.
        ws = wb['Absensi'] if 'Absensi' in wb.sheetnames else wb.worksheets[0]
        if ws.title.casefold() in ('contoh', 'petunjuk'):
            raise XlsxError('Sheet Absensi tidak ditemukan.')
        ws.reset_dimensions()  # do not trust attacker-controlled dimension metadata
        iterator = ws.iter_rows()
        header = next(iterator, ())
        titles = tuple(str(c.value or '').strip() for c in header)
        if titles[:8] != HEADERS or any(titles[8:]):
            raise XlsxError('Kolom wajib berurutan: ' + ', '.join(HEADERS) + '.')
        rows = []
        for number, cells in enumerate(iterator, 2):
            if number > MAX_ROWS + 1:
                raise XlsxError('Maksimal 5000 baris data.')
            if not any(c.value is not None and str(c.value).strip() for c in cells):
                continue
            errors = []
            if any(c.value is not None and str(c.value).strip() for c in cells[8:]):
                errors.append('Data di luar delapan kolom tidak diperbolehkan.')
            values = {}
            for index, field in enumerate(FIELDS):
                cell = cells[index] if index < len(cells) else None
                value = cell.value if cell else None
                if cell and cell.data_type == 'f':
                    errors.append('Formula tidak diperbolehkan.')
                if field in ('nik', 'nip', 'phone') and value is not None and not isinstance(value, str):
                    errors.append(f'{HEADERS[index]} harus disimpan sebagai teks, bukan angka (risiko kehilangan digit).')
                values[field] = str(value).strip() if value is not None else ''
            rows.append({'row': number, **values, '_errors': errors})
        return rows
    except XlsxError:
        raise
    except (ValueError, TypeError, KeyError, IndexError, BadZipFile, ElementTree.ParseError, OSError, InvalidFileException) as exc:
        raise XlsxError('File bukan workbook XLSX yang valid.') from exc
    finally:
        if wb is not None:
            wb.close()


def _name(value):
    return ' '.join(value.split()).casefold()


def preview(event, source_rows, dry_run):
    participants = list(Participant.objects.filter(event=event))
    by_nik = {p.nik: p for p in participants}
    by_name = defaultdict(set)
    for p in participants:
        by_name[_name(p.full_name)].add(p.nik)
    attendances = {a.participant_id: a for a in Attendance.objects.filter(event=event)}
    seen_niks, seen_names = {}, {}
    result = {'rows': [], 'created': 0, 'updated': 0, 'errors': [], 'dry_run': dry_run}
    for source in source_rows:
        row = {k: v for k, v in source.items() if k != '_errors'}
        messages = list(source['_errors'])
        p = by_nik.get(row['nik'])
        a = attendances.get(p.pk) if p else None
        row.update(participant_id=str(p.pk) if p else None, attendance_id=str(a.pk) if a else None, action=None)
        candidate = Participant(event=event)
        for field in FIELDS[:-1]:
            try:
                Participant._meta.get_field(field).clean(row[field], candidate)
            except ValidationError as exc:
                messages.extend(f'{HEADERS[FIELDS.index(field)]}: {m}' for m in exc.messages)
        if row['status'] not in Attendance.Status.values:
            messages.append('Status harus hadir atau tidak_hadir.')
        nik, name = row['nik'], _name(row['full_name'])
        if nik in seen_niks:
            messages.append(f'NIK duplikat dengan baris {seen_niks[nik]}.')
        if name and name in seen_names:
            messages.append(f'Nama duplikat/ambigu dengan baris {seen_names[name]}.')
        seen_niks[nik], seen_names[name] = row['row'], row['row']
        if p and _name(p.full_name) != name:
            messages.append('NIK sudah terdaftar pada kegiatan ini dengan nama berbeda.')
        if name and by_name[name] - {nik}:
            messages.append('Nama sudah terdaftar pada kegiatan ini dengan NIK lain; identitas ambigu.')
        if messages:
            result['errors'].append({'row': row['row'], 'message': '; '.join(messages)})
        else:
            row['action'] = 'updated' if a else 'created'
            result[row['action']] += 1
        result['rows'].append(row)
    return result


def apply_rows(event, result):
    """Caller owns atomic transaction and event lock; validate before any writes."""
    for row in result['rows']:
        defaults = {field: row[field] for field in FIELDS[:-1] if row[field] != ''}
        defaults.pop('nik', None)
        if row['nip']:
            defaults['is_asn'] = True
        participant, _ = Participant.objects.update_or_create(event=event, nik=row['nik'], defaults=defaults)
        attendance, _ = Attendance.objects.update_or_create(event=event, participant=participant, defaults={'status': row['status']})
        row['participant_id'] = str(participant.pk)
        row['attendance_id'] = str(attendance.pk)
