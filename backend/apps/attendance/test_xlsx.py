from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.events.models import Event
from apps.participants.models import Participant
from .models import Attendance

HEADERS = ['NIK', 'NIP', 'Nama Lengkap', 'Instansi', 'Jabatan', 'No HP', 'Email', 'Status']
VALID = ['0123456789012345', '012345678901234567', 'Nama Tes', 'Diskominfo', 'Staf', '08123456789', 'nama@example.com', 'hadir']


class AttendanceXlsxTests(TestCase):
    def setUp(self):
        now = timezone.now()
        self.event = Event.objects.create(title='Excel Test', start_date=now, end_date=now)
        self.other = Event.objects.create(title='Other', start_date=now, end_date=now)
        self.user = get_user_model().objects.create_user(username='xlsx-admin', role='admin')
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.base = f'/api/events/{self.event.pk}/attendance/'

    def upload(self, rows=None, dry_run=None, name='attendance.xlsx', workbook=None):
        wb = workbook or Workbook()
        if workbook is None:
            wb.active.title = 'Absensi'
            wb.active.append(HEADERS)
            for row in rows or []:
                wb.active.append(row)
        out = BytesIO()
        wb.save(out)
        data = {'file': SimpleUploadedFile(name, out.getvalue())}
        if dry_run is not None:
            data['dry_run'] = dry_run
        return self.client.post(self.base + 'import-xlsx/', data, format='multipart')

    def test_template_empty_data_and_separate_example(self):
        response = self.client.get(self.base + 'template-xlsx/')
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb['Absensi'].max_row, 1)
        self.assertEqual([c.value for c in wb['Absensi'][1]], HEADERS)
        self.assertIn('Petunjuk', wb.sheetnames)
        self.assertIn('Contoh', wb.sheetnames)
        self.assertEqual(wb['Contoh']['B2'].data_type, 's')
        self.assertEqual(wb['Contoh']['A3'].data_type, 's')
        response = self.upload(workbook=wb, dry_run='false')
        self.assertEqual(response.status_code, 400)
        self.assertIn('kosong', response.data['errors'][0]['message'])
        self.assertEqual(response.data['created'], 0)
        self.assertFalse(Participant.objects.exists())

    def test_default_preview_has_contract_and_no_writes_or_fake_ids(self):
        response = self.upload([VALID])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(set(response.data), {'rows', 'created', 'updated', 'errors', 'dry_run'})
        self.assertTrue(response.data['dry_run'])
        self.assertEqual(response.data['created'], 1)
        self.assertEqual(response.data['rows'][0]['nik'], VALID[0])
        self.assertIsNone(response.data['rows'][0]['participant_id'])
        self.assertFalse(Participant.objects.exists())
        self.assertFalse(Attendance.objects.exists())

    def test_apply_and_repeat_are_event_scoped_idempotent_no_certificates(self):
        Participant.objects.create(event=self.other, nik=VALID[0], full_name=VALID[2])
        with patch('apps.certificates.services.generate_certificates_for_event') as generate:
            response = self.upload([VALID], 'false')
            self.assertEqual(response.status_code, 200)
            generate.assert_not_called()
        participant = Participant.objects.get(event=self.event)
        attendance = Attendance.objects.get(event=self.event)
        self.assertEqual(participant.nip, VALID[1])
        self.assertTrue(participant.is_asn)
        self.assertEqual(response.data['rows'][0]['participant_id'], str(participant.pk))
        second = self.upload([VALID], 'false')
        self.assertEqual(second.data['created'], 0)
        self.assertEqual(second.data['updated'], 1)
        self.assertEqual(Attendance.objects.get(event=self.event).pk, attendance.pk)
        self.assertEqual(Attendance.objects.count(), 1)

    def test_upsert_preserves_blank_optional_fields_and_changes_status(self):
        self.upload([VALID], 'false')
        row = [VALID[0], '', VALID[2], '', '', '', '', 'tidak_hadir']
        result = self.upload([row], 'false')
        self.assertEqual(result.status_code, 200)
        participant = Participant.objects.get(event=self.event)
        self.assertEqual(participant.phone, VALID[5])
        self.assertEqual(participant.nip, VALID[1])
        self.assertEqual(Attendance.objects.get().status, 'tidak_hadir')

    def test_all_invalid_rows_report_errors_without_partial_import(self):
        bad = ['1', '123', '', '', '', '', 'bad-email', 'wrong']
        response = self.upload([VALID, bad, bad], 'false')
        self.assertEqual(response.status_code, 400)
        self.assertEqual({e['row'] for e in response.data['errors']}, {3, 4})
        self.assertFalse(Participant.objects.exists())

    def test_duplicate_nik_and_name_in_file_are_rejected(self):
        row = [*VALID]
        row[0] = '1123456789012345'
        response = self.upload([VALID, VALID, row], 'false')
        self.assertEqual(response.status_code, 400)
        self.assertTrue(response.data['errors'])
        self.assertFalse(Attendance.objects.exists())

    def test_existing_nik_name_mismatch_is_rejected(self):
        Participant.objects.create(event=self.event, nik=VALID[0], full_name='Different Name')
        self.assertEqual(self.upload([VALID], 'false').status_code, 400)
        self.assertFalse(Attendance.objects.exists())

    def test_existing_name_with_other_nik_is_ambiguous(self):
        Participant.objects.create(event=self.event, nik='1123456789012345', full_name=VALID[2])
        self.assertEqual(self.upload([VALID], 'false').status_code, 400)

    def test_numeric_identity_is_rejected_instead_of_float_conversion(self):
        for column in (0, 1):
            with self.subTest(column=column):
                row = [*VALID]
                row[column] = int(row[column])
                self.assertEqual(self.upload([row]).status_code, 400)

    def test_nik_only_and_at_least_one_identity_required(self):
        row = [*VALID]
        row[1] = ''
        self.assertEqual(self.upload([row], 'false').status_code, 200)
        self.assertFalse(Participant.objects.get().is_asn)
        row[0] = ''
        self.assertEqual(self.upload([row]).status_code, 400)

    def test_formula_rejected_even_on_other_sheet(self):
        wb = Workbook()
        wb.active.title = 'Absensi'
        wb.active.append(HEADERS)
        wb.active.append(VALID)
        wb.create_sheet('Other')['A1'] = '=1+1'
        self.assertEqual(self.upload(workbook=wb).status_code, 400)

    def test_formula_in_data_rejected(self):
        row = [*VALID]
        row[2] = '=HYPERLINK("https://example.com")'
        self.assertEqual(self.upload([row]).status_code, 400)

    def test_export_text_safe_and_all_statuses_scoped(self):
        row = [*VALID]
        row[2] = '=Not a formula'
        participant = Participant.objects.create(event=self.event, nik=row[0], nip=row[1], full_name=row[2])
        Attendance.objects.create(event=self.event, participant=participant, status='tidak_hadir')
        other = Participant.objects.create(event=self.other, nik='1123456789012345', full_name='Other')
        Attendance.objects.create(event=self.other, participant=other)
        response = self.client.get(self.base + 'export-xlsx/')
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        ws = wb['Absensi']
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws['A2'].value, row[0])
        self.assertEqual(ws['B2'].value, row[1])
        self.assertEqual(ws['C2'].value, row[2])
        self.assertEqual(ws['C2'].data_type, 's')
        self.assertEqual(ws['H2'].value, 'tidak_hadir')

    def test_wrong_extension_malformed_zip_and_missing_file(self):
        self.assertEqual(self.upload([VALID], name='bad.csv').status_code, 400)
        response = self.client.post(self.base + 'import-xlsx/', {'file': SimpleUploadedFile('bad.xlsx', b'not zip')}, format='multipart')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.client.post(self.base + 'import-xlsx/', {}, format='multipart').status_code, 400)

    def test_dry_run_invalid_value_rejected(self):
        self.assertEqual(self.upload([VALID], 'no').status_code, 400)

    def test_row_limit_and_zip_bomb_limits(self):
        self.assertEqual(self.upload([VALID] * 5001).status_code, 400)
        out = BytesIO()
        with ZipFile(out, 'w', ZIP_DEFLATED) as z:
            z.writestr('xl/worksheets/sheet1.xml', b'0' * (26 * 1024 * 1024))
        response = self.client.post(self.base + 'import-xlsx/', {'file': SimpleUploadedFile('bomb.xlsx', out.getvalue())}, format='multipart')
        self.assertEqual(response.status_code, 400)

    def test_file_size_limit(self):
        response = self.client.post(self.base + 'import-xlsx/', {'file': SimpleUploadedFile('large.xlsx', b'0' * (5 * 1024 * 1024 + 1))}, format='multipart')
        self.assertEqual(response.status_code, 400)

    def test_headers_and_extra_data_columns_rejected(self):
        wb = Workbook()
        wb.active.append(['Bad Header'])
        self.assertEqual(self.upload(workbook=wb).status_code, 400)
        self.assertEqual(self.upload([[*VALID, 'Unexpected']]).status_code, 400)

    def test_permissions_all_endpoints_and_unknown_event(self):
        for role in ('operator', 'admin', 'superadmin'):
            self.user.role = role
            self.user.save()
            for action in ('template-xlsx/', 'export-xlsx/'):
                expected = 403 if role == 'operator' else 200
                self.assertEqual(self.client.get(self.base + action).status_code, expected)
            self.assertEqual(self.upload([VALID]).status_code, 403 if role == 'operator' else 200)
        self.client.force_authenticate(None)
        self.assertIn(self.client.get(self.base + 'template-xlsx/').status_code, (401, 403))
        self.assertIn(self.upload([VALID]).status_code, (401, 403))
        self.client.force_authenticate(self.user)
        self.assertEqual(self.client.get('/api/events/00000000-0000-0000-0000-000000000000/attendance/export-xlsx/').status_code, 404)

    def test_corrupt_xlsx_package_and_row_indices_return_validation_errors(self):
        out = BytesIO()
        with ZipFile(out, 'w') as z:
            z.writestr('unrelated.txt', 'not an XLSX package')
        response = self.client.post(self.base + 'import-xlsx/', {'file': SimpleUploadedFile('invalid.xlsx', out.getvalue())}, format='multipart')
        self.assertEqual(response.status_code, 400)
        self.assertIn('errors', response.data)
        wb = Workbook()
        wb.active.append(HEADERS)
        wb.active.append(VALID)
        original = BytesIO()
        wb.save(original)
        out = BytesIO()
        with ZipFile(original) as src, ZipFile(out, 'w') as dst:
            for name in src.namelist():
                data = src.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'<row r="2">', b'<row r="invalid">')
                dst.writestr(name, data)
        response = self.client.post(self.base + 'import-xlsx/', {'file': SimpleUploadedFile('invalid.xlsx', out.getvalue())}, format='multipart')
        self.assertEqual(response.status_code, 400)

    def test_model_field_lengths_email_and_invalid_nip_are_validated(self):
        for index, value in ((0, '1' * 17), (1, '1' * 17), (2, 'N' * 201), (3, 'I' * 201), (4, 'J' * 151), (5, '1' * 21), (6, 'not-email')):
            row = [*VALID]
            row[index] = value
            with self.subTest(column=index):
                response = self.upload([row], 'false')
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.data['errors'][0]['row'], 2)
        self.assertFalse(Attendance.objects.exists())

    def test_boundary_5000_rows_preview(self):
        rows = [[f'{i:016d}', '', f'Person {i}', '', '', '', '', 'hadir'] for i in range(5000)]
        response = self.upload(rows)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['created'], 5000)
        self.assertFalse(Participant.objects.exists())

    def test_database_failure_rolls_back_all_rows(self):
        row = [*VALID]
        row[0], row[1], row[2] = '1123456789012345', '112345678901234567', 'Second Name'
        original = Attendance.objects.update_or_create
        calls = 0

        def fail_second(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                from django.db import IntegrityError
                raise IntegrityError('race')
            return original(*args, **kwargs)

        with patch.object(Attendance.objects, 'update_or_create', side_effect=fail_second):
            response = self.upload([VALID, row], 'false')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Participant.objects.exists())
        self.assertFalse(Attendance.objects.exists())
